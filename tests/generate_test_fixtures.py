import os
import io
import json
import csv
import zipfile
import xml.etree.ElementTree as ET
from PIL import Image, ImageDraw

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "fixtures")
os.makedirs(FIXTURES_DIR, exist_ok=True)

def create_sample_png_bytes(color=(18, 197, 110), text="Sample Image"):
    img = Image.new("RGB", (200, 100), color=color)
    draw = ImageDraw.Draw(img)
    draw.text((20, 40), text, fill=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()

def generate_ieee_pdf():
    """1. 12-page IEEE PDF with 13 images."""
    import fitz
    pdf_path = os.path.join(FIXTURES_DIR, "ieee_paper_12pages_13images.pdf")
    doc = fitz.open()

    img_bytes = create_sample_png_bytes()

    for p_idx in range(12):
        page = doc.new_page()
        page_num = p_idx + 1
        page.insert_text(
            (50, 50),
            f"IEEE TRANSACTIONS ON COMPUTATIONAL AUDITING, VOL. 42, NO. 3\n"
            f"Page {page_num}: Evaluation of Distributed Ledger Verification Protocols.\n\n"
            f"Section {page_num}: Empirical observations regarding throughput in cluster environment.\n"
            f"Paragraph on page {page_num} documents experimental telemetry and convergence rates."
        )

        # Insert 1 image on pages 1-11, and 2 images on page 12 -> total 13 images
        images_to_insert = 2 if page_num == 12 else 1
        for img_i in range(images_to_insert):
            rect = fitz.Rect(50, 150 + (img_i * 120), 250, 250 + (img_i * 120))
            page.insert_image(rect, stream=img_bytes)

    # Page 12 has References
    p12 = doc[11]
    p12.insert_text(
        (50, 450),
        "\nReferences\n"
        "[1] Smith, J. and Doe, A. (2021). Scalable Quantum Verification. Nature Quantum.\n"
        "[2] Brown, C. et al. (2023). Deterministic Evidence Provenance. IEEE S&P."
    )

    doc.save(pdf_path)
    doc.close()
    print("Created ieee_paper_12pages_13images.pdf")

def generate_pptx():
    """2. PPTX with images, tables, speaker notes, and claim on slide 3."""
    import pptx
    from pptx.util import Inches, Pt

    pptx_path = os.path.join(FIXTURES_DIR, "presentation_with_notes.pptx")
    prs = pptx.Presentation()

    img_bytes = create_sample_png_bytes((59, 130, 246), "Slide Chart")
    img_stream = io.BytesIO(img_bytes)

    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    slide1.shapes.title.text = "Q3 Performance Review"
    slide1.placeholders[1].text = "Technical Operations and Quantum Computing Benchmarks"

    # Slide 2: Table
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Infrastructure Milestones"
    table_shape = slide2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Uptime"
    table.cell(1, 1).text = "99.995%"
    table.cell(2, 0).text = "Global Nodes"
    table.cell(2, 1).text = "450"

    # Slide 3: Claim slide with speaker notes and image
    slide3 = prs.slides.add_slide(prs.slide_layouts[5])
    slide3.shapes.title.text = "Quantum Processing Milestones"
    tx_box = slide3.shapes.add_textbox(Inches(1), Inches(1.8), Inches(8), Inches(1.5))
    tf = tx_box.text_frame
    tf.text = "Quantum supremacy benchmark was completed in 200 seconds across 53 qubits."
    slide3.shapes.add_picture(img_stream, Inches(1), Inches(3.5), width=Inches(3))
    
    # Speaker notes on Slide 3
    slide3.notes_slide.notes_text_frame.text = "Note for presenter: Highlight the 200 seconds speedup vs supercomputer baseline."

    # Slide 4: Concluding
    slide4 = prs.slides.add_slide(prs.slide_layouts[5])
    slide4.shapes.title.text = "Next Steps"
    tx4 = slide4.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    tx4.text_frame.text = "Transitioning to fault-tolerant logical qubit validation in Q4."

    prs.save(pptx_path)
    print("Created presentation_with_notes.pptx")

def generate_xlsx():
    """3. XLSX with several sheets; Sheet 'Q3' row 12 has specific data."""
    import openpyxl
    xlsx_path = os.path.join(FIXTURES_DIR, "multisheet_financials.xlsx")
    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws1 = wb.active
    ws1.title = "Summary"
    ws1.append(["Metric", "Status"])
    ws1.append(["Overall Audit", "Passed"])

    # Sheet 2: Q3 (row 12 has target claim)
    ws2 = wb.create_sheet(title="Q3")
    ws2.append(["Item", "Yield", "Target", "Status"]) # Row 1 headers
    for r in range(2, 12):
        ws2.append([f"Batch {r-1}", f"{80 + r}%", "85%", "Standard"])
    # Row 12: Target Claim
    ws2.append(["BC-902 clinical yield", "94.8%", "90%", "Optimal"])

    # Sheet 3: Projections
    ws3 = wb.create_sheet(title="Projections")
    ws3.append(["Year", "Forecast"])
    ws3.append(["2025", "$250M"])

    wb.save(xlsx_path)
    print("Created multisheet_financials.xlsx")

def generate_json():
    """4. Nested JSON with specific value."""
    json_path = os.path.join(FIXTURES_DIR, "nested_infrastructure.json")
    data = {
        "metadata": {
            "environment": "production",
            "version": "2.4.0"
        },
        "infrastructure": {
            "datacenter": {
                "cluster": {
                    "primary_region": "us-east-1",
                    "node_count": 128,
                    "hypervisor": "KVM"
                }
            }
        }
    }
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    print("Created nested_infrastructure.json")

def generate_csv():
    """5. CSV with columns and rows."""
    csv_path = os.path.join(FIXTURES_DIR, "products_catalog.csv")
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(["sku", "product_name", "inventory_count", "unit_price"])
        writer.writerow(["SKU-101", "Titanium Sensor", "450", "$129.99"])
        writer.writerow(["SKU-102", "Optic Transceiver", "120", "$89.50"])
        writer.writerow(["SKU-103", "Precision Calibrator", "85", "$340.00"])
    print("Created products_catalog.csv")

def generate_docx():
    """6. DOCX with embedded images and tables."""
    import docx
    from docx.shared import Inches

    docx_path = os.path.join(FIXTURES_DIR, "document_with_images.docx")
    doc = docx.Document()
    doc.add_heading("Cloud Architecture Protocol Specification", level=1)
    doc.add_paragraph("This standard specifies container hardening and permission bounds.")
    
    # Table
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "Security Layer"
    t.cell(0, 1).text = "Policy Rule"
    t.cell(1, 0).text = "Network Proxy"
    t.cell(1, 1).text = "Mutual TLS enforced on port 8443"

    # Embedded image
    img_bytes = create_sample_png_bytes((168, 85, 247), "Architecture Diagram")
    img_stream = io.BytesIO(img_bytes)
    doc.add_picture(img_stream, width=Inches(3))

    doc.add_heading("References", level=2)
    doc.add_paragraph("[1] NIST Special Publication 800-190, Application Container Security Guide.")

    doc.save(docx_path)
    print("Created document_with_images.docx")

def generate_bib():
    """7. .bib file."""
    bib_path = os.path.join(FIXTURES_DIR, "academic_references.bib")
    content = """@article{smith2021quantum,
  author = {Smith, John and Doe, Jane},
  title = {Scalable Quantum Verification on Hybrid Cloud Infrastructure},
  journal = {Nature Quantum Information},
  year = {2021},
  volume = {7},
  pages = {42--50},
  doi = {10.1038/s41534-021-00001-x}
}

@article{miller2022distributed,
  author = {Miller, Robert and Chen, Wei},
  title = {Automated Verification of Neural Representations},
  journal = {Journal of AI Research},
  year = {2022},
  doi = {10.1613/jair.1.13401}
}"""
    with open(bib_path, "w", encoding="utf-8") as f:
        f.write(content)
    print("Created academic_references.bib")

def generate_scanned_pdf():
    """8. Image-only scanned PDF (no extractable text characters)."""
    import fitz
    pdf_path = os.path.join(FIXTURES_DIR, "scanned_image_only.pdf")
    doc = fitz.open()

    img_bytes = create_sample_png_bytes((75, 85, 99), "Scanned Contract Page")

    # Page 1: only an image, no text
    page1 = doc.new_page()
    rect1 = fitz.Rect(50, 50, 500, 700)
    page1.insert_image(rect1, stream=img_bytes)

    # Page 2: only an image, no text
    page2 = doc.new_page()
    rect2 = fitz.Rect(50, 50, 500, 700)
    page2.insert_image(rect2, stream=img_bytes)

    doc.save(pdf_path)
    doc.close()
    print("Created scanned_image_only.pdf")

if __name__ == "__main__":
    generate_ieee_pdf()
    generate_pptx()
    generate_xlsx()
    generate_json()
    generate_csv()
    generate_docx()
    generate_bib()
    generate_scanned_pdf()
    print("All fixtures generated successfully.")
